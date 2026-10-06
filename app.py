from pathlib import Path
import hashlib
import json
import re
import uuid
from datetime import datetime

import chromadb
import ollama
import streamlit as st
from docx import Document
from pypdf import PdfReader


# ---------- Settings ----------
CHAT_MODEL = "llama3.2"
EMBED_MODEL = "nomic-embed-text"

DOCUMENTS_DIR = Path("documents")
UPLOADS_DIR = DOCUMENTS_DIR / "uploads"
DATA_DIR = Path("data")
CHROMA_DIR = DATA_DIR / "chroma"
CHAT_HISTORY_FILE = DATA_DIR / "chat_history.json"

CHUNK_SIZE = 800
CHUNK_OVERLAP = 120
RESULT_COUNT = 5
SUPPORTED_TYPES = {".txt", ".pdf", ".docx"}

DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
CHROMA_DIR.mkdir(parents=True, exist_ok=True)
DATA_DIR.mkdir(parents=True, exist_ok=True)


# ---------- Chroma database ----------
@st.cache_resource
def get_chroma_client():
    return chromadb.PersistentClient(path=str(CHROMA_DIR))


client = get_chroma_client()
collection = client.get_or_create_collection(name="local_ai_documents")


# ---------- Saved chat history ----------
def load_chat_history():
    if not CHAT_HISTORY_FILE.exists():
        return []

    try:
        history = json.loads(CHAT_HISTORY_FILE.read_text(encoding="utf-8"))
        return history if isinstance(history, list) else []
    except (json.JSONDecodeError, OSError):
        return []


def save_chat_history():
    CHAT_HISTORY_FILE.write_text(
        json.dumps(st.session_state.chat_history, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def save_current_chat():
    messages = st.session_state.messages
    if not messages:
        return

    chat_id = st.session_state.active_chat_id

    if chat_id is None:
        first_question = next(
            (
                message["content"]
                for message in messages
                if message["role"] == "user"
            ),
            "New chat",
        )
        title = first_question.strip()
        if len(title) > 45:
            title = title[:42] + "..."

        chat_id = str(uuid.uuid4())
        st.session_state.active_chat_id = chat_id
        st.session_state.chat_history.append(
            {
                "id": chat_id,
                "title": title or "New chat",
                "updated_at": datetime.now().isoformat(timespec="seconds"),
                "messages": [],
            }
        )

    for chat in st.session_state.chat_history:
        if chat["id"] == chat_id:
            chat["messages"] = messages
            chat["updated_at"] = datetime.now().isoformat(timespec="seconds")
            break

    save_chat_history()


if "chat_history" not in st.session_state:
    st.session_state.chat_history = load_chat_history()

if "active_chat_id" not in st.session_state:
    st.session_state.active_chat_id = None

if "messages" not in st.session_state:
    st.session_state.messages = []


# ---------- Read and split documents ----------
def split_into_chunks(text):
    text = " ".join(text.split())
    chunks = []
    start = 0

    while start < len(text):
        chunk = text[start : start + CHUNK_SIZE].strip()
        if chunk:
            chunks.append(chunk)
        start += CHUNK_SIZE - CHUNK_OVERLAP

    return chunks


def read_document(path):
    sections = []
    extension = path.suffix.lower()

    if extension == ".txt":
        text = path.read_text(encoding="utf-8", errors="ignore")
        sections.append((text, None))

    elif extension == ".pdf":
        reader = PdfReader(str(path))

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                sections.append((text, page_number))

    elif extension == ".docx":
        document = Document(str(path))
        parts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                )
                if row_text:
                    parts.append(row_text)

        if parts:
            sections.append(("\n".join(parts), None))

    return sections


def make_id(source, page_number, chunk_number, text):
    unique_text = f"{source}|{page_number}|{chunk_number}|{text}"
    return hashlib.sha256(unique_text.encode("utf-8")).hexdigest()


# ---------- Build the document index ----------
def rebuild_index():
    files = [
        path
        for path in DOCUMENTS_DIR.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_TYPES
    ]

    records = {}

    for path in files:
        source = path.relative_to(DOCUMENTS_DIR).as_posix()

        try:
            sections = read_document(path)
        except Exception as error:
            st.warning(f"Could not read {source}: {error}")
            continue

        for section_text, page_number in sections:
            for chunk_number, chunk in enumerate(split_into_chunks(section_text)):
                chunk_id = make_id(source, page_number, chunk_number, chunk)

                metadata = {"source": source}
                if page_number is not None:
                    metadata["page"] = page_number

                records[chunk_id] = {
                    "text": chunk,
                    "metadata": metadata,
                }

    if not records:
        return 0

    ids = list(records.keys())
    texts = [records[item_id]["text"] for item_id in ids]
    metadatas = [records[item_id]["metadata"] for item_id in ids]

    # Create all embeddings before replacing the current index.
    embeddings = []
    for text in texts:
        result = ollama.embeddings(model=EMBED_MODEL, prompt=text)
        embeddings.append(result["embedding"])

    old_records = collection.get(include=[])
    old_ids = old_records.get("ids", [])
    if old_ids:
        collection.delete(ids=old_ids)

    collection.add(
        ids=ids,
        documents=texts,
        metadatas=metadatas,
        embeddings=embeddings,
    )

    return len(ids)


# ---------- Find relevant passages ----------
STOP_WORDS = {
    "a", "an", "and", "are", "as", "at", "be", "by", "did", "do", "does",
    "for", "from", "how", "i", "in", "is", "it", "me", "of", "on", "or",
    "our", "please", "that", "the", "their", "this", "to", "was", "what",
    "when", "where", "which", "who", "why", "with", "would",
}


def question_terms(question):
    words = re.findall(r"\b[a-zA-Z0-9]+\b", question.lower())
    return {word for word in words if len(word) > 2 and word not in STOP_WORDS}


def find_passages(question):
    total = collection.count()
    if total == 0:
        return []

    query_embedding = ollama.embeddings(
        model=EMBED_MODEL,
        prompt=question,
    )["embedding"]

    semantic = collection.query(
        query_embeddings=[query_embedding],
        n_results=min(15, total),
        include=["documents", "metadatas", "distances"],
    )

    candidates = {}

    for index, text in enumerate(semantic["documents"][0]):
        metadata = semantic["metadatas"][0][index]
        distance = semantic["distances"][0][index]
        key = f"{metadata.get('source')}|{metadata.get('page')}|{text}"

        candidates[key] = {
            "text": text,
            "metadata": metadata,
            "distance": distance,
        }

    all_records = collection.get(include=["documents", "metadatas"])
    terms = question_terms(question)

    for text, metadata in zip(
        all_records.get("documents", []),
        all_records.get("metadatas", []),
    ):
        text_words = set(re.findall(r"\b[a-zA-Z0-9]+\b", text.lower()))
        overlap = len(terms & text_words)

        if overlap:
            key = f"{metadata.get('source')}|{metadata.get('page')}|{text}"
            if key not in candidates:
                candidates[key] = {
                    "text": text,
                    "metadata": metadata,
                    "distance": 999,
                }

    def rank(item):
        text_words = set(re.findall(r"\b[a-zA-Z0-9]+\b", item["text"].lower()))
        overlap = len(terms & text_words)
        return (-overlap, item["distance"])

    return sorted(candidates.values(), key=rank)[:RESULT_COUNT]


def direct_answer_from_passages(question, passages):
    terms = question_terms(question)
    best_sentence = ""
    best_score = 0

    for passage in passages:
        sentences = re.split(r"(?<=[.!?])\s+", passage["text"])

        for sentence in sentences:
            sentence_words = set(
                re.findall(r"\b[a-zA-Z0-9]+\b", sentence.lower())
            )
            score = len(terms & sentence_words)

            if score > best_score:
                best_score = score
                best_sentence = sentence.strip()

    if terms and best_score >= max(1, len(terms) // 2):
        return best_sentence

    return None


# ---------- Streamlit page ----------
st.set_page_config(page_title="Local AI Knowledge Retrieval")
st.title("Local AI Knowledge Retrieval")

with st.form("upload_form", clear_on_submit=True):
    uploaded_files = st.file_uploader(
        "Upload PDF, TXT, or Word files",
        type=["pdf", "txt", "docx"],
        accept_multiple_files=True,
    )
    save_uploads = st.form_submit_button("Save uploaded files")

if save_uploads and uploaded_files:
    for uploaded_file in uploaded_files:
        save_path = UPLOADS_DIR / Path(uploaded_file.name).name
        save_path.write_bytes(uploaded_file.getvalue())

    st.success("Files saved. Click **Rebuild document index** next.")

if st.button("Rebuild document index"):
    try:
        with st.spinner("Reading and indexing your documents..."):
            indexed_count = rebuild_index()

        if indexed_count:
            st.success(f"Done. Indexed {indexed_count} text sections.")
        else:
            st.warning("No readable text found. Add a PDF, TXT, or DOCX file first.")

    except Exception as error:
        st.error(
            "Indexing failed. Make sure Ollama is running and "
            "nomic-embed-text is installed."
        )
        st.caption(str(error))


# ---------- Sidebar ----------
with st.sidebar:
    st.header("Manage uploaded files")

    uploaded_paths = sorted(
        [
            path
            for path in UPLOADS_DIR.iterdir()
            if path.is_file() and path.suffix.lower() in SUPPORTED_TYPES
        ],
        key=lambda path: path.name.lower(),
    )

    if uploaded_paths:
        selected_name = st.selectbox(
            "Choose a file to delete",
            [path.name for path in uploaded_paths],
        )

        if st.button("Delete selected file"):
            selected_path = UPLOADS_DIR / selected_name
            source = selected_path.relative_to(DOCUMENTS_DIR).as_posix()

            collection.delete(where={"source": source})
            selected_path.unlink()

            st.success(f"Deleted {selected_name}")
            st.rerun()
    else:
        st.caption("No uploaded files.")

    st.divider()
    st.header("Chat history")

    if st.button("New chat", use_container_width=True):
        st.session_state.active_chat_id = None
        st.session_state.messages = []
        st.rerun()

    history = st.session_state.chat_history
    sorted_history = sorted(
        history,
        key=lambda item: item.get("updated_at", ""),
        reverse=True,
    )
    history_ids = [chat["id"] for chat in sorted_history]
    options = [None] + history_ids
    titles = {chat["id"]: chat.get("title", "Saved chat") for chat in history}

    active_id = st.session_state.active_chat_id
    selector_key = f"chat_selector_{active_id or 'new'}"
    selected_index = options.index(active_id) if active_id in options else 0

    selected_chat_id = st.selectbox(
        "Open a saved chat",
        options,
        index=selected_index,
        format_func=lambda chat_id: (
            "New chat" if chat_id is None else titles.get(chat_id, "Saved chat")
        ),
        key=selector_key,
    )

    if selected_chat_id != active_id:
        st.session_state.active_chat_id = selected_chat_id

        if selected_chat_id is None:
            st.session_state.messages = []
        else:
            selected_chat = next(
                chat for chat in history if chat["id"] == selected_chat_id
            )
            st.session_state.messages = selected_chat.get("messages", [])

        st.rerun()

    if st.session_state.active_chat_id is not None:
        if st.button("Delete current chat", use_container_width=True):
            current_id = st.session_state.active_chat_id
            st.session_state.chat_history = [
                chat
                for chat in st.session_state.chat_history
                if chat["id"] != current_id
            ]
            save_chat_history()
            st.session_state.active_chat_id = None
            st.session_state.messages = []
            st.rerun()

    st.caption("Chat history is saved locally on this computer.")


# ---------- Show current conversation ----------
for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

question = st.chat_input("Ask a question about your documents")

if question:
    st.session_state.messages.append(
        {"role": "user", "content": question}
    )
    save_current_chat()

    with st.chat_message("user"):
        st.markdown(question)

    with st.chat_message("assistant"):
        try:
            passages = find_passages(question)

            if not passages:
                answer = "I couldn’t find that in the documents."
            else:
                context_parts = []
                sources = []

                for passage in passages:
                    metadata = passage["metadata"]
                    source = metadata.get("source", "Unknown file")
                    page = metadata.get("page")

                    citation = f"{source} (page {page})" if page else source
                    context_parts.append(
                        f"Source: {citation}\n{passage['text']}"
                    )
                    sources.append(citation)

                context = "\n\n".join(context_parts)

                response = ollama.chat(
                    model=CHAT_MODEL,
                    messages=[
                        {
                            "role": "system",
                            "content": (
                                "Answer using only the provided passages. "
                                "If a passage contains the answer, answer it "
                                "clearly and briefly. Never say the answer is "
                                "missing when it appears in a passage. "
                                "If none of the passages contain the answer, "
                                "say: I couldn’t find that in the documents."
                            ),
                        },
                        {
                            "role": "user",
                            "content": (
                                f"Question: {question}\n\n"
                                f"Passages:\n{context}"
                            ),
                        },
                    ],
                    options={"temperature": 0},
                )

                answer = response["message"]["content"].strip()

                refusal_phrases = (
                    "couldn’t find",
                    "couldn't find",
                    "cannot find",
                    "not found in",
                    "not mentioned",
                    "does not contain",
                )

                if any(phrase in answer.lower() for phrase in refusal_phrases):
                    direct_answer = direct_answer_from_passages(
                        question, passages
                    )
                    if direct_answer:
                        answer = direct_answer

                answer += "\n\nSources: " + ", ".join(
                    dict.fromkeys(sources)
                )

                with st.expander("Retrieved passages"):
                    st.text(context)

            st.markdown(answer)

        except Exception as error:
            answer = (
                "I couldn’t connect to Ollama. Check that Ollama is running "
                "and that the models are installed."
            )
            st.error(answer)
            st.caption(str(error))

    st.session_state.messages.append(
        {"role": "assistant", "content": answer}
    )
    save_current_chat()


# ---------- Download latest answer ----------
latest_answer = next(
    (
        message["content"]
        for message in reversed(st.session_state.messages)
        if message["role"] == "assistant"
    ),
    None,
)

if latest_answer:
    st.download_button(
        "Download latest answer",
        data=latest_answer,
        file_name="answer.txt",
        mime="text/plain",
        key=f"download_answer_{st.session_state.active_chat_id or 'new'}_"
            f"{len(st.session_state.messages)}",
    )